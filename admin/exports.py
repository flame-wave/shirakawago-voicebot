"""書き出し（Excel・CSV）。

  ・質問の一覧 … 「登録されている質問」から。会議の資料や、職員どうしの確かめに使う
  ・利用状況のレポート … 「利用状況」から。期間を決めて、集計と記録を1つのExcelにまとめる

画面で見ている数字と同じになるよう、集計は log_stats と page_stats の決まりをそのまま使う。
Excel は見出しを太字・色付きにし、列の幅を中身に合わせる（開いてすぐ読めるように）。
CSV は Excel でそのまま開けるよう、UTF-8（BOM付き）にする。
"""
import io
from collections import Counter
from datetime import datetime, timedelta, timezone

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import log_stats as L

JST = timezone(timedelta(hours=9))
LANG_NAMES = {"ja": "日本語", "en": "英語", "zh": "中国語", "ko": "韓国語",
              "es": "スペイン語", "fr": "フランス語"}
TRANSLATED = [("en", "英語"), ("zh", "中国語"), ("ko", "韓国語"), ("es", "スペイン語"), ("fr", "フランス語")]
VOTE_NAMES = [("good", "〇 よかった"), ("ok", "△ ふつう"), ("bad", "× いまいち")]
SOURCE_NAMES = {"faq": "質問回答集", "ai": "AI", "none": "職員へ回った"}
TR_STATES = {"none": "訳なし", "changed": "訳し直しが要る", "machine": "要確認", "ok": "確かめ済み"}


# ---------------------------------------------------------------- 共通
def to_excel(sheets) -> bytes:
    """{シート名: DataFrame} を1つのExcelにする。見出しを目立たせ、列の幅を合わせる。"""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name[:31], index=False)
            ws = writer.sheets[name[:31]]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2C5F2D")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            for c, column in enumerate(frame.columns, start=1):
                values = [str(column)] + [str(v) for v in frame[column].head(200)]
                longest = max((sum(2 if ord(ch) > 0x2E80 else 1 for ch in v) for v in values), default=8)
                ws.column_dimensions[get_column_letter(c)].width = min(max(8, longest + 2), 60)
                if longest > 60:
                    for row in ws.iter_rows(min_row=2, min_col=c, max_col=c):
                        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    return buf.getvalue()


def to_csv(frame) -> bytes:
    """Excel でそのまま開ける CSV（UTF-8・BOM付き）。"""
    return frame.to_csv(index=False).encode("utf-8-sig")


def _jst(row):
    at = L._when(row)
    return at.astimezone(JST).strftime("%Y-%m-%d %H:%M") if at else ""


# ---------------------------------------------------------------- 質問の一覧
def questions_frame(book, items, state_of, place_label, categories):
    """質問の一覧の表。state_of(item) は状態の名前、place_label(place) は設置場所の書き方。"""
    status = book.translation_status([l for l, _ in TRANSLATED])
    rows = []
    for it in items:
        tr = book.translations_of(it["id"])
        row = {
            "管理用の名前": it["id"],
            "状態": state_of(it),
            "設置場所": place_label(it["place"]),
            "分類": categories.get(it["category"], "その他"),
            "聞き方（日本語）": "\n".join(it["questions"]),
            "回答（日本語）": it["answer"],
            "読み上げ用の短い回答": it["short"],
        }
        for lang, name in TRANSLATED:
            row[f"回答（{name}）"] = (tr.get(lang) or {}).get("answer", "")
        per = status.get(it["id"], {})
        need = [name for lang, name in TRANSLATED if per.get(lang) in ("none", "changed")]
        check = [name for lang, name in TRANSLATED if per.get(lang) == "machine"]
        row.update({
            "翻訳の状態": ("訳が要る：" + "・".join(need)) if need
                         else ("要確認：" + "・".join(check)) if check else "そろっている",
            "写真": it["photo"],
            "参考ページ": it["link"],
            "よくある質問の札": it["chip"],
            "案内の始まり": it["show_from"][:10],
            "案内の終わり": it["show_until"][:10],
            "最終更新": it["updated"],
        })
        rows.append(row)
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- 利用状況のレポート
def usage_report(rows, ai_rows, survey_rows, label_of, period_name) -> bytes:
    """利用状況のレポート（Excel）。rows は質問の記録、ai_rows はAIの記録、survey_rows はアンケート。

    label_of(faq_id) は、質問回答集の項目を人が読める名前にする（page_stats._label と同じもの）。
    """
    s = L.summarize(rows)
    by = s["by_source"]
    talk = sum(1 for r in rows if str(r.get("faq_id") or "").startswith("talk:"))
    votes = Counter(str(r.get("vote")) for r in survey_rows)
    n_votes = sum(votes[v] for v, _ in VOTE_NAMES)
    pct = lambda n, d: f"{n / d:.0%}" if d else "—"

    summary = [
        ("期間", period_name),
        ("作成した日時", datetime.now(JST).strftime("%Y-%m-%d %H:%M")),
        ("質問の総数", s["total"]),
        ("その場で答えられた割合", pct(s["answered"], s["total"])),
        ("質問回答集で答えた（会話を含む）", by.get("faq", 0)),
        ("うち、あいさつなどの会話", talk),
        ("AIが答えた", by.get("ai", 0)),
        ("職員へ回った（答えられなかった）", by.get("none", 0)),
        ("アンケートの回答数", n_votes),
    ] + [(f"アンケート {name}", f"{votes[v]} 件（{pct(votes[v], n_votes)}）") for v, name in VOTE_NAMES]
    sheets = {"概要": pd.DataFrame(summary, columns=["項目", "値"])}

    # 日ごと（案内所ごとの件数）
    clients = L.clients_in(rows)
    days = {}
    for r in rows:
        at = L._when(r)
        if at is None:
            continue
        d = at.astimezone(JST).strftime("%Y-%m-%d")
        days.setdefault(d, Counter())[str(r.get("client") or "不明")] += 1
    daily = [{"日付": d, **{c: days[d][c] for c in clients}, "計": sum(days[d].values())}
             for d in sorted(days)]
    sheets["日ごと"] = pd.DataFrame(daily, columns=["日付", *clients, "計"])

    # 案内所ごと
    per_client = L.by_client(rows)
    sheets["案内所ごと"] = pd.DataFrame([
        {"案内所": c, "質問の数": v["total"], "その場で答えられた割合": pct(v["answered"], v["total"]),
         "AIが答えた": v["by_source"].get("ai", 0), "職員へ回った": v["by_source"].get("none", 0)}
        for c, v in per_client.items()])

    # 言語ごと
    sheets["言語ごと"] = pd.DataFrame(
        sorted(([LANG_NAMES.get(k, k), n] for k, n in s["by_lang"].items()), key=lambda x: -x[1]),
        columns=["言語", "質問の数"])

    # よく聞かれた質問・答えられなかった質問
    sheets["よく聞かれた質問"] = pd.DataFrame(
        [(i + 1, label_of(fid), fid, n) for i, (fid, n) in enumerate(s["top_faq"])],
        columns=["順位", "答えた内容", "管理用の名前", "回数"])
    where = {}
    for r in rows:
        if str(r.get("source") or "none") == "none" and str(r.get("recognized") or "").strip():
            where.setdefault(str(r["recognized"]).strip(), set()).add(str(r.get("client") or "不明"))
    sheets["答えられなかった質問"] = pd.DataFrame(
        [(text, n, "・".join(sorted(where.get(text, [])))) for text, n in s["unmatched"]],
        columns=["聞かれた言葉", "回数", "案内所"])

    # AIが答えた質問（雑談の返事は別に数える）
    answered = sorted((r for r in ai_rows if r.get("result") == "answered"),
                      key=lambda r: str(r.get("at") or ""), reverse=True)
    sheets["AIが答えた質問"] = pd.DataFrame(
        [(_jst(r), LANG_NAMES.get(r.get("lang", ""), r.get("lang", "")), r.get("place") or "",
          r.get("question", ""), r.get("answer", "")) for r in answered],
        columns=["日時（日本時間）", "言語", "案内所", "質問", "AIの答え"])

    # アンケート（案内所ごと・言語ごと）
    def survey_table(key, names=None):
        groups = {}
        for r in survey_rows:
            g = str(r.get(key) or "不明")
            groups.setdefault(names.get(g, g) if names else g, Counter())[str(r.get("vote"))] += 1
        out = []
        for g, c in sorted(groups.items(), key=lambda x: -sum(x[1].values())):
            total = sum(c[v] for v, _ in VOTE_NAMES)
            out.append({"": g, **{name: c[v] for v, name in VOTE_NAMES}, "計": total,
                        "〇の割合": pct(c["good"], total)})
        return pd.DataFrame(out, columns=["", *[n for _, n in VOTE_NAMES], "計", "〇の割合"])
    sheets["アンケート（案内所ごと）"] = survey_table("client")
    sheets["アンケート（言語ごと）"] = survey_table("lang", LANG_NAMES)

    # 記録そのもの（いちばん後ろ）
    sheets["記録"] = pd.DataFrame(
        [(_jst(r), r.get("client", ""), r.get("place") or "",
          LANG_NAMES.get(r.get("lang", ""), r.get("lang", "")), r.get("recognized", ""),
          SOURCE_NAMES.get(str(r.get("source") or "none"), ""), r.get("faq_id") or "")
         for r in sorted(rows, key=lambda r: str(r.get("at") or ""))],
        columns=["日時（日本時間）", "端末", "現在地", "言語", "聞かれた言葉", "答え方", "答えた質問"])
    return to_excel(sheets)
