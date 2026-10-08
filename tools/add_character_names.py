"""「キャラクター」シートに、言語ごとの名前の欄を足す。

案内アプリを英語や中国語にしても、キャラクターの名前だけ日本語のままだと、
読めない方には「何を選ぶボタンか」が分からない。
言語ごとの名前を持てるようにする。空の欄は日本語の名前のまま出る。

欄は表の右端に足す（既存の欄の位置を変えると、管理画面や
build_faq.py が読む場所がずれるおそれがあるため）。
すでに欄があれば何もしない。

使い方:
    python tools/add_character_names.py
"""
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PATH = ROOT / "tools" / "faq_master.xlsx"
SHEET = "キャラクター"

# 欄の名前と言語。build_faq.py の CHARACTER_NAME_COLUMNS と同じ並び。
COLUMNS = [
    ("名前（英語）", "en"),
    ("名前（中国語）", "zh"),
    ("名前（韓国語）", "ko"),
    ("名前（スペイン語）", "es"),
    ("名前（フランス語）", "fr"),
]

# いま登録されているキャラクターの名前（どれも仮）。
# 英語・スペイン語・フランス語はローマ字、中国語は意味の通る呼び名、韓国語は読みをハングルで。
NAMES = {
    "navi": ("Navi", "小导", "나비", "Navi", "Navi"),
    "main": ("Yui", "结衣", "유이", "Yui", "Yui"),
    "yuru": ("Gassho-kun", "合掌君", "갓쇼 군", "Gassho-kun", "Gassho-kun"),
    "yane": ("Yane-kun", "屋顶君", "야네 군", "Yane-kun", "Yane-kun"),
    "gou": ("Go-chan", "小乡", "고 짱", "Go-chan", "Go-chan"),
}


def main():
    wb = load_workbook(PATH)
    if SHEET not in wb.sheetnames:
        print(f"「{SHEET}」シートがありません。")
        return
    ws = wb[SHEET]
    header = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    if COLUMNS[0][0] in header:
        print("言語ごとの名前の欄は、すでにあります。中身は触りません。")
        return

    id_col = header.get("ID")
    start = ws.max_column + 1
    for i, (name, _) in enumerate(COLUMNS):
        cell = ws.cell(row=1, column=start + i, value=name)
        cell.font = Font(name="Yu Gothic", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="EDF4E6")
        ws.column_dimensions[get_column_letter(start + i)].width = 16

    filled = 0
    for r in range(2, ws.max_row + 1):
        cid = str(ws.cell(row=r, column=id_col).value or "").strip() if id_col else ""
        names = NAMES.get(cid)
        if not names:
            continue
        for i, value in enumerate(names):
            cell = ws.cell(row=r, column=start + i, value=value)
            cell.font = Font(name="Yu Gothic", size=10)
        filled += 1

    wb.save(PATH)
    print(f"言語ごとの名前の欄を {len(COLUMNS)} つ足し、{filled} 体に名前を入れました。")
    print("名前は管理画面の「キャラクター」→「表で細かく直す」から変えられます。")


if __name__ == "__main__":
    main()
