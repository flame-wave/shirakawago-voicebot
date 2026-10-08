"""「会話」シートを作る（「こんにちは」「ありがとう」などへの返事）。

質問回答集は「案内」のための表なので、あいさつのような短い会話は入れない。
キャラクターごとに話し方を変えられるよう、キャラクターの側に返事を持たせる。

  種類         … あいさつ／お礼／名前／元気／さようなら（言い方は案内アプリが知っている）
                  それ以外の名前を書くと、自分で足した種類になる（「好きな食べ物」など）。
  キャラクター … キャラクターのID（main など）か「共通」。
                  そのキャラクターの行が無い種類は「共通」の返事を使う。
  聞き方       … 自分で足した種類で、どう言われたらこの返事をするか（1行に1つ。どの言語でもよい）。
                  あいさつなど5つの種類では空でよい（足すと、その言い方でも当たる）。
  返事         … 日本語の返事。{名前} と書くと、そのキャラクターの名前に置き換わる。
                  同じ種類・同じキャラクターの行を何行か書くと、その中から選んで答える。
  返事（英語）など … 各言語の返事。管理画面で保存すると自動で翻訳される。

使い方:
    python tools/add_talk.py              （無ければ作る）
    python tools/add_talk.py --recreate   （作り直す。書いた返事は消えるので注意）
"""
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PATH = ROOT / "tools" / "faq_master.xlsx"

SHEET = "会話"
HEADERS = ["種類", "キャラクター", "聞き方", "返事", "返事（英語）", "返事（中国語）", "返事（韓国語）",
           "返事（スペイン語）", "返事（フランス語）", "備考"]
WIDTHS = [12, 12, 30, 46, 40, 34, 34, 40, 40, 20]

FOOD = "\n".join(["好きな食べ物", "何が好き", "好物", "なに食べたい", "favorite food",
                  "what's your favorite food", "what is your favorite food",
                  "what do you like to eat", "喜欢吃什么", "最喜欢的食物", "좋아하는 음식",
                  "좋아하는 음식이 뭐예요", "comida favorita", "cuál es tu comida favorita",
                  "plat préféré", "quel est ton plat préféré"])
HOBBY = "\n".join(["趣味", "趣味は何", "休みの日は何してる", "好きなこと", "hobby", "hobbies",
                   "what's your hobby", "what are your hobbies", "what do you do for fun",
                   "爱好", "你的爱好是什么", "兴趣", "취미", "취미가 뭐예요", "pasatiempo",
                   "cuál es tu pasatiempo", "passe-temps", "loisirs", "quel est ton passe-temps"])
PRAISE = "\n".join(["かわいい", "かわいいね", "かわいいですね", "すごいね", "賢いね", "cute",
                    "you are cute", "so cute", "smart", "可爱", "好可爱", "귀여워", "귀엽다",
                    "귀여워요", "linda", "qué linda", "mignonne", "trop mignon"])

# (種類, 聞き方, 日本語, 英語, 中国語, 韓国語, スペイン語, フランス語)
ROWS = [
    ("あいさつ", "",
     "こんにちは！白川郷へようこそ。わたしは{名前}です。知りたいことを聞いてくださいね。",
     "Hello! Welcome to Shirakawa-go. I'm {名前}. Feel free to ask me anything.",
     "你好！欢迎来到白川乡。我是{名前}，有什么想知道的请问我吧。",
     "안녕하세요! 시라카와고에 오신 것을 환영합니다. 저는 {名前}입니다. 궁금한 것을 물어보세요.",
     "¡Hola! Bienvenido a Shirakawa-go. Soy {名前}. Pregúntame lo que quieras.",
     "Bonjour ! Bienvenue à Shirakawa-go. Je suis {名前}. N'hésitez pas à me poser vos questions."),
    ("あいさつ", "",
     "こんにちは！{名前}です。白川郷のことなら、なんでも聞いてください。",
     "Hi there! I'm {名前}. Ask me anything about Shirakawa-go.",
     "你好！我是{名前}。关于白川乡的事情，尽管问我吧。",
     "안녕하세요! {名前}입니다. 시라카와고에 대해 무엇이든 물어보세요.",
     "¡Hola! Soy {名前}. Pregúntame lo que quieras sobre Shirakawa-go.",
     "Salut ! Je suis {名前}. Posez-moi toutes vos questions sur Shirakawa-go."),
    ("あいさつ", "",
     "ようこそ、白川郷へ！合掌造りの里を、ゆっくり楽しんでいってくださいね。",
     "Welcome to Shirakawa-go! Take your time and enjoy the gassho-style village.",
     "欢迎来到白川乡！请慢慢欣赏合掌造村落吧。",
     "시라카와고에 오신 것을 환영해요! 갓쇼즈쿠리 마을을 천천히 즐겨 주세요.",
     "¡Bienvenido a Shirakawa-go! Disfrute con calma del pueblo de casas gassho.",
     "Bienvenue à Shirakawa-go ! Prenez le temps de profiter du village aux maisons gassho."),
    ("お礼", "",
     "どういたしまして。白川郷を楽しんでくださいね。",
     "You're welcome! Enjoy your time in Shirakawa-go.",
     "不客气！祝您在白川乡玩得愉快。",
     "천만에요! 시라카와고에서 즐거운 시간 보내세요.",
     "¡De nada! Disfrute de Shirakawa-go.",
     "Je vous en prie ! Profitez bien de Shirakawa-go."),
    ("お礼", "",
     "お役に立ててうれしいです。ほかにも気になることがあれば、聞いてくださいね。",
     "I'm glad I could help. Let me know if there's anything else you'd like to know.",
     "很高兴能帮上忙。如果还有其他想知道的，请随时问我。",
     "도움이 되어 기뻐요. 더 궁금한 것이 있으면 물어보세요.",
     "Me alegra haber ayudado. Si tiene otra pregunta, no dude en hacerla.",
     "Je suis ravie d'avoir pu aider. N'hésitez pas si vous avez d'autres questions."),
    ("名前", "",
     "わたしは{名前}です。白川郷の案内をしています。",
     "I'm {名前}, your guide to Shirakawa-go.",
     "我是{名前}，负责介绍白川乡。",
     "저는 {名前}입니다. 시라카와고를 안내하고 있어요.",
     "Soy {名前}, su guía de Shirakawa-go.",
     "Je suis {名前}, votre guide à Shirakawa-go."),
    ("名前", "",
     "{名前}といいます。どうぞよろしくお願いします！",
     "My name is {名前}. Nice to meet you!",
     "我叫{名前}，请多关照！",
     "{名前}라고 해요. 잘 부탁드려요!",
     "Me llamo {名前}. ¡Encantada de conocerle!",
     "Je m'appelle {名前}. Enchantée !"),
    ("元気", "",
     "はい、元気です！何かお手伝いできることはありますか。",
     "I'm doing great, thank you! How can I help you?",
     "我很好，谢谢！有什么可以帮您的吗？",
     "네, 잘 지내요! 무엇을 도와드릴까요?",
     "¡Muy bien, gracias! ¿En qué puedo ayudarle?",
     "Très bien, merci ! Comment puis-je vous aider ?"),
    ("元気", "",
     "とっても元気です！今日も白川郷はいいところですよ。",
     "I'm full of energy! Shirakawa-go is lovely today, too.",
     "我精神很好！今天的白川乡也很美哦。",
     "아주 잘 지내요! 오늘도 시라카와고는 정말 좋은 곳이에요.",
     "¡Con mucha energía! Shirakawa-go también está precioso hoy.",
     "En pleine forme ! Shirakawa-go est magnifique aujourd'hui aussi."),
    ("さようなら", "",
     "ありがとうございました。どうぞお気をつけて、白川郷を楽しんでください。",
     "Thank you for visiting! Take care and enjoy Shirakawa-go.",
     "谢谢您！路上小心，祝您在白川乡玩得开心。",
     "감사합니다! 조심히 다니시고 시라카와고를 즐기세요.",
     "¡Gracias! Cuídese y disfrute de Shirakawa-go.",
     "Merci ! Prenez soin de vous et profitez de Shirakawa-go."),
    ("さようなら", "",
     "またいつでも来てくださいね。いってらっしゃい！",
     "Please come back anytime. Have a great trip!",
     "欢迎随时再来。一路顺风！",
     "언제든지 또 오세요. 좋은 여행 되세요!",
     "Vuelva cuando quiera. ¡Buen viaje!",
     "Revenez quand vous voulez. Bon voyage !"),
    ("好きな食べ物", FOOD,
     "わたしは朴葉味噌が大好きです。香ばしくて、ご飯がすすみますよ。",
     "I love hoba miso! It's fragrant and goes perfectly with rice.",
     "我最喜欢朴叶味噌了！香喷喷的，特别下饭。",
     "저는 호바미소를 정말 좋아해요. 고소해서 밥이 술술 넘어가요.",
     "¡Me encanta el hoba miso! Es aromático y va perfecto con arroz.",
     "J'adore le hoba miso ! Il est parfumé et se marie parfaitement avec le riz."),
    ("好きな食べ物", FOOD,
     "五平餅が好きです！甘いたれの香りがたまりません。",
     "I like gohei-mochi! The sweet sauce smells wonderful.",
     "我喜欢五平饼！甜甜的酱汁香气真让人受不了。",
     "고헤이모치를 좋아해요! 달콤한 소스 향이 최고예요.",
     "¡Me gusta el gohei-mochi! La salsa dulce huele de maravilla.",
     "J'aime le gohei-mochi ! La sauce sucrée sent délicieusement bon."),
    ("趣味", HOBBY,
     "合掌造りの屋根をながめるのが好きです。季節ごとに表情が変わるんですよ。",
     "I love looking at the thatched gassho roofs. They change with every season.",
     "我喜欢眺望合掌造的茅草屋顶，每个季节的样子都不一样哦。",
     "갓쇼즈쿠리 지붕을 바라보는 걸 좋아해요. 계절마다 모습이 달라져요.",
     "Me encanta mirar los techos de paja de las casas gassho. Cambian con cada estación.",
     "J'adore contempler les toits de chaume des maisons gassho. Ils changent à chaque saison."),
    ("ほめられた", PRAISE,
     "ありがとうございます！そう言ってもらえるとうれしいです。",
     "Thank you! That makes me happy.",
     "谢谢！听您这么说我很开心。",
     "고마워요! 그렇게 말해 주시니 기뻐요.",
     "¡Gracias! Me alegra mucho oír eso.",
     "Merci ! Ça me fait très plaisir."),
    ("ほめられた", PRAISE,
     "えへへ、照れちゃいます。白川郷も素敵なところですよ！",
     "Hehe, you're making me blush. Shirakawa-go is lovely too!",
     "嘿嘿，有点不好意思呢。白川乡也是个很棒的地方哦！",
     "헤헤, 쑥스러워요. 시라카와고도 멋진 곳이에요!",
     "Jeje, me da vergüenza. ¡Shirakawa-go también es precioso!",
     "Hihi, vous me faites rougir. Shirakawa-go est aussi un endroit merveilleux !"),
]

GUIDE = (
    "※ 「こんにちは」「ありがとう」などへの返事です。質問回答集で見つからなかったときに使います。"
    "　種類は あいさつ／お礼／名前／元気／さようなら のほか、自分で名前を付けて足せます（「聞き方」に言い方を1行に1つ）。"
    "　キャラクターには ID（main など）か「共通」を書きます。そのキャラクターの行が無い種類は「共通」の返事になります。"
    "　同じ種類の行を何行か書くと、その中から選んで答えます。{名前} は、そのキャラクターの名前に置き換わります。"
    "　管理画面の「キャラクター」から書けます（保存すると自動で翻訳されます）。"
)


def build(wb):
    at = wb.sheetnames.index("キャラクター") + 1 if "キャラクター" in wb.sheetnames else None
    ws = wb.create_sheet(SHEET, at)
    for c, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(name="Yu Gothic", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="EDF4E6")
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]
    ws.freeze_panes = "D2"
    for r, (kind, phrases, *replies) in enumerate(ROWS, start=2):
        values = [kind, "共通", phrases, *replies, ""]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name="Yu Gothic", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=c >= 3)
    row = len(ROWS) + 3
    guide = ws.cell(row=row, column=1, value=GUIDE)
    guide.font = Font(name="Yu Gothic", size=9, color="5A6B58")
    guide.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 3, end_column=len(HEADERS))


def main():
    wb = load_workbook(PATH)
    if SHEET in wb.sheetnames:
        if "--recreate" not in sys.argv:
            print(f"「{SHEET}」シートはすでにあります。中身は触りません（作り直すときは --recreate）。")
            return
        del wb[SHEET]
    build(wb)
    wb.save(PATH)
    kinds = len({r[0] for r in ROWS})
    print(f"「{SHEET}」シートを作り、共通の返事を {kinds} 種類・{len(ROWS)} 通り入れました。")


if __name__ == "__main__":
    main()
