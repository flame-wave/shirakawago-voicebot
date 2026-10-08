"""
回答の読み上げ音声を VOICEVOX でまとめて作る。

ブラウザの合成音声はキャラクターの見た目と合わないため、
よく聞かれる回答だけでも、決まった声で用意しておくためのツール。
作った音声ファイルがある回答は、案内アプリがそちらを優先して再生する。

────────────────────────────────────────────────
準備:
  1. VOICEVOX（https://voicevox.hiroshiba.jp/）を入れて起動する
     （アプリを起動しておけば、裏で音声合成の窓口が動く）
  2. 話者の一覧を見る
       python tools/make_voice.py --list
  3. 作る
       python tools/make_voice.py --speaker 3
       python tools/make_voice.py --speaker 3 --only toilet bus_next

  「読み方」に言葉を足したあと、その言葉を含む音声だけを作り直す場合:
       python tools/make_voice.py --fix-readings            （何を作り直すか見るだけ）
       python tools/make_voice.py --fix-readings --go       （作り直して mp3 にする）
       python tools/make_voice.py --fix-readings --words 朴葉味噌 --go

  別の声も用意して、案内アプリ側で選べるようにする場合:
       python tools/make_voice.py --speaker 8 --set zundamon --label "ずんだもん"
     → assets/audio/zundamon/ に作られ、「音声セット」シートに登録される
────────────────────────────────────────────────

作った音声は assets/audio/ に置かれ、Excelの「音声ファイル名」列に
ファイル名が書き込まれる。あとは管理画面から書き出せば案内アプリに反映される。

【作ったあとに必ずやること】
VOICEVOXが書き出すWAVは1件0.5MBほどある。観光客のスマートフォンは
回答のたびにこれを読み込むので、mp3に変換しておく（10分の1になる）。

    python tools/convert_audio.py

【使用上の注意】
VOICEVOXで作った音声を公開・設置で使う場合は、
話者ごとの利用規約に従い、クレジット表記（例: VOICEVOX:四国めたん）が必要。
"""

import argparse
import sys
from pathlib import Path

import requests
from openpyxl import load_workbook

ENGINE = "http://127.0.0.1:50021"
ROOT = Path(__file__).resolve().parent.parent

# 作った音声を登録する先。案内アプリはここから使う声を選べる。
VOICE_SHEET = "音声セット"
VOICE_HEADERS = ["名前", "フォルダ", "話者番号", "速さ", "備考"]

# 読み上げに使う文。短い回答があればそちらを優先する。
SHORT_COL = "読み上げ用の短い回答"
ANSWER_COL = "回答"
AUDIO_COL = "音声ファイル名"
ID_COL = "ID"
ENABLED_COL = "有効"

# 読み間違いを直す表。案内アプリのブラウザ読み上げと同じシートを使う。
READING_SHEET = "読み方"


def check_engine(engine: str) -> None:
    try:
        res = requests.get(f"{engine}/version", timeout=5)
        print(f"VOICEVOX に接続しました（版: {res.text.strip()}）")
    except requests.RequestException:
        print(
            "VOICEVOX に接続できません。\n"
            "  VOICEVOXアプリを起動してから、もう一度実行してください。\n"
            f"  （接続先: {engine}）"
        )
        sys.exit(1)


def list_speakers(engine: str) -> None:
    speakers = requests.get(f"{engine}/speakers", timeout=10).json()
    print("\n話者の一覧（--speaker に番号を指定する）\n")
    for speaker in speakers:
        for style in speaker["styles"]:
            print(f"  {style['id']:>3}  {speaker['name']}（{style['name']}）")
    print("\n※ 利用規約は話者ごとに異なります。設置で使う前に必ず確認してください。")


def load_readings(wb) -> list:
    """「読み方」シートを読む。無ければ空。

    VOICEVOXは「荻町」を「はぎまち」、「朴葉味噌」を「ぼくようみそ」と読む。
    合成に回す文だけカタカナに差し替える（Excelの回答は書き換えない）。

    長い語から先に差し替える。「八幡神社」を先に直してしまうと
    「白川八幡神社」が半端に差し替わる。
    """
    if READING_SHEET not in wb.sheetnames:
        return []
    ws = wb[READING_SHEET]
    header = {cell(ws, 1, c): c for c in range(1, ws.max_column + 1)}
    if "表記" not in header or "読み" not in header:
        return []

    out = []
    for r in range(2, ws.max_row + 1):
        surface = cell(ws, r, header["表記"])
        reading = cell(ws, r, header["読み"])
        if not surface or not reading or surface.startswith("※"):
            continue
        if "有効" in header and cell(ws, r, header["有効"]).upper()                 in ("FALSE", "0", "×", "NO"):
            continue
        out.append((surface, reading))
    out.sort(key=lambda pair: -len(pair[0]))
    return out


def apply_readings(text: str, readings: list) -> str:
    for surface, reading in readings:
        text = text.replace(surface, reading)
    return text


def synthesize(engine: str, text: str, speaker: int, speed: float) -> bytes:
    """文章から音声（WAV）を作る。"""
    query = requests.post(
        f"{engine}/audio_query",
        params={"text": text, "speaker": speaker},
        timeout=30,
    )
    query.raise_for_status()
    data = query.json()
    data["speedScale"] = speed

    wav = requests.post(
        f"{engine}/synthesis",
        params={"speaker": speaker},
        json=data,
        timeout=180,
    )
    wav.raise_for_status()
    return wav.content


def cell(ws, row: int, col: int) -> str:
    value = ws.cell(row, col).value
    return "" if value is None else str(value).strip()


def register_set(wb, folder: str, label: str, speaker: int, speed: float) -> None:
    """作った音声を「音声セット」シートに登録する。

    案内アプリは、この一覧から使う声を選べるようにしている。
    同じフォルダの行が既にあれば書き換える。
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if VOICE_SHEET in wb.sheetnames:
        ws = wb[VOICE_SHEET]
    else:
        ws = wb.create_sheet(VOICE_SHEET)
        for c, name in enumerate(VOICE_HEADERS, start=1):
            head = ws.cell(row=1, column=c, value=name)
            head.font = Font(name="Yu Gothic", size=10, bold=True)
            head.fill = PatternFill("solid", fgColor="EDF4E6")
        ws.freeze_panes = "A2"
        for c, width in enumerate([24, 18, 12, 10, 40], start=1):
            ws.column_dimensions[get_column_letter(c)].width = width

    row = None
    for r in range(2, ws.max_row + 1):
        if cell(ws, r, 2) == folder and cell(ws, r, 1):
            row = r
            break
    if row is None:
        row = ws.max_row + 1
        while row > 2 and not cell(ws, row - 1, 1):
            row -= 1

    values = [label, folder, speaker, speed,
              "VOICEVOXで作成。設置で使う場合はクレジット表記が必要"]
    for c, value in enumerate(values, start=1):
        target = ws.cell(row=row, column=c, value=value)
        target.font = Font(name="Yu Gothic", size=10)
        target.alignment = Alignment(vertical="top", wrap_text=(c == 5))


def voice_sets(wb) -> list:
    """「音声セット」シートの一覧 [(名前, フォルダ, 話者番号, 速さ)]。話者番号の無い行は飛ばす。"""
    if VOICE_SHEET not in wb.sheetnames:
        return []
    ws = wb[VOICE_SHEET]
    header = {cell(ws, 1, c): c for c in range(1, ws.max_column + 1)}
    out = []
    for r in range(2, ws.max_row + 1):
        name = cell(ws, r, header.get("名前", 1))
        if not name or name.startswith("※"):
            continue
        try:
            speaker = int(float(cell(ws, r, header["話者番号"])))
        except (KeyError, ValueError):
            continue
        try:
            speed = float(cell(ws, r, header.get("速さ", 0)) or 1.0) if "速さ" in header else 1.0
        except ValueError:
            speed = 1.0
        out.append((name, cell(ws, r, header["フォルダ"]) if "フォルダ" in header else "",
                    speaker, speed))
    return out


def speaker_names(engine: str) -> dict:
    names = {}
    for speaker in requests.get(f"{engine}/speakers", timeout=10).json():
        for style in speaker["styles"]:
            names[style["id"]] = f"{speaker['name']}（{style['name']}）"
    return names


def fix_readings(args) -> None:
    """読み方の表にある言葉を含む回答の音声だけを、すべての音声セットで作り直す。

    「読み方」に言葉を足しても、すでに作った音声は変わらない（VOICEVOXで作り済みのため）。
    全部を作り直すと時間がかかるので、その言葉を含む回答だけを作り直す。
    --go を付けないときは、何を作り直すかを見せるだけで何も書き換えない。
    """
    check_engine(args.engine)
    wb = load_workbook(Path(args.source))
    all_readings = load_readings(wb)   # 合成に回すときは、表のすべてを当てる
    readings = all_readings
    if args.words:
        readings = [(s, r) for s, r in readings if s in args.words]
    if not readings:
        print("当てる読み方がありません（「読み方」シートか --words を確かめてください）。")
        return

    ws = wb["FAQ"]
    header = {cell(ws, 1, c): c for c in range(1, ws.max_column + 1)}
    targets = []
    for r in range(2, ws.max_row + 1):
        fid = cell(ws, r, header[ID_COL])
        audio = cell(ws, r, header[AUDIO_COL]) if AUDIO_COL in header else ""
        if not fid or not audio:
            continue
        text = (cell(ws, r, header[SHORT_COL]) if SHORT_COL in header else "") \
            or cell(ws, r, header[ANSWER_COL])
        words = [s for s, _ in readings if s in text]
        if words:
            targets.append((fid, audio, text, words))
    if not targets:
        print("読み方の表にある言葉を含む、用意した音声はありません。")
        return

    names = speaker_names(args.engine)
    sets = voice_sets(wb)
    print(f"\n作り直す回答: {len(targets)} 件")
    for fid, audio, _, words in targets:
        print(f"  {fid}（{audio}）… {'・'.join(words)}")
    print("\n声のセット:")
    for name, folder, speaker, speed in sets:
        print(f"  {name} … 話者番号 {speaker} = {names.get(speaker, '不明')}"
              f"／速さ {speed}／assets/audio/{folder or ''}")
    if not args.go:
        print("\nまだ何も書き換えていません。話者番号が、いまの音声と同じ声か確かめてから、"
              "--go を付けてもう一度実行してください。")
        return

    import tempfile
    from convert_audio import convert, find_ffmpeg
    ffmpeg = find_ffmpeg()
    made, failed = [], []
    for name, folder, speaker, speed in sets:
        out_dir = Path(args.out) / folder if folder else Path(args.out)
        for fid, audio, text, _ in targets:
            target = out_dir / audio
            if not target.exists():
                continue   # このセットには、この回答の音声がもともと無い
            try:
                wav = synthesize(args.engine, apply_readings(text, all_readings), speaker, speed)
            except requests.RequestException as e:
                failed.append(f"{name}/{audio}: {e}")
                continue
            if target.suffix.lower() == ".mp3":
                if not ffmpeg:
                    failed.append(f"{name}/{audio}: ffmpeg が無いため mp3 にできません")
                    continue
                with tempfile.TemporaryDirectory() as tmp:
                    src = Path(tmp) / "voice.wav"
                    src.write_bytes(wav)
                    if not convert(ffmpeg, src, target):
                        failed.append(f"{name}/{audio}: mp3 にできませんでした")
                        continue
            else:
                target.write_bytes(wav)
            made.append(target.relative_to(ROOT))
            print(f"  作り直しました: {target.relative_to(ROOT)}")
    print(f"\n作り直した音声: {len(made)} 件")
    for f in failed:
        print("  × " + f)
    if made:
        print("\nこのあと:")
        print("  1. 作り直したファイルを、ロリポップの同じ場所（assets/audio/…）へ FTP で上げる")
        print("  2. git push する（GitHub の控えも新しくする）")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", default=str(ROOT / "tools" / "faq_master.xlsx"))
    ap.add_argument("--out", default=str(ROOT / "assets" / "audio"))
    ap.add_argument("--engine", default=ENGINE)
    ap.add_argument("--speaker", type=int, help="話者の番号（--list で確認）")
    ap.add_argument("--speed", type=float, default=1.0, help="話す速さ（既定 1.0）")
    ap.add_argument("--only", nargs="*", help="作る項目のID（省略すると全件）")
    ap.add_argument("--overwrite", action="store_true", help="既にある音声も作り直す")
    ap.add_argument("--set", default="",
                    help="置き場所（assets/audio の下のフォルダ名）。省略すると直下に作る")
    ap.add_argument("--label", default="",
                    help="案内アプリに出すこの声の名前（省略すると話者名）")
    ap.add_argument("--list", action="store_true", help="話者の一覧を表示して終わる")
    ap.add_argument("--fix-readings", action="store_true",
                    help="「読み方」の表にある言葉を含む回答の音声だけを、すべての声のセットで作り直す")
    ap.add_argument("--words", nargs="*",
                    help="--fix-readings で対象にする言葉（省略すると「読み方」の表のすべて）")
    ap.add_argument("--go", action="store_true",
                    help="--fix-readings で実際に作り直す（付けないと、作り直すものを見せるだけ）")
    args = ap.parse_args()

    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    if args.fix_readings:
        fix_readings(args)
        return

    check_engine(args.engine)
    if args.list:
        list_speakers(args.engine)
        return
    if args.speaker is None:
        print("--speaker を指定してください（--list で番号を確認できます）")
        sys.exit(1)

    out_dir = Path(args.out) / args.set if args.set else Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    path = Path(args.source)
    wb = load_workbook(path)
    ws = wb["FAQ"]
    header = {cell(ws, 1, c): c for c in range(1, ws.max_column + 1)}
    for name in (ID_COL, ANSWER_COL):
        if name not in header:
            print(f"FAQシートに「{name}」列がありません")
            sys.exit(1)
    if AUDIO_COL not in header:
        print(f"FAQシートに「{AUDIO_COL}」列がありません")
        sys.exit(1)

    readings = load_readings(wb)
    if readings:
        print(f"読み方の直しを {len(readings)} 語あてます"
              f"（「{READING_SHEET}」シート）")
    else:
        print(f"「{READING_SHEET}」シートが無いため、読み方の直しは当てません。")
        print("  地名や料理名が読み間違えられる場合は "
              "python tools/add_readings.py で作れます。")

    made, skipped, failed = 0, 0, []
    for r in range(2, ws.max_row + 1):
        fid = cell(ws, r, header[ID_COL])
        if not fid:
            continue
        if args.only and fid not in args.only:
            continue
        if ENABLED_COL in header and cell(ws, r, header[ENABLED_COL]).upper() \
                not in ("TRUE", "1", "○", "YES", ""):
            continue

        # 読み上げ用の短い回答があればそれを使う（長い回答は聞き取りにくいため）
        text = ""
        if SHORT_COL in header:
            text = cell(ws, r, header[SHORT_COL])
        if not text:
            text = cell(ws, r, header[ANSWER_COL])
        if not text:
            continue

        filename = cell(ws, r, header[AUDIO_COL]) or f"{fid}.wav"
        target = out_dir / filename
        if target.exists() and not args.overwrite:
            skipped += 1
            continue

        spoken = apply_readings(text, readings)
        print(f"  作成中: {fid} → {filename}")
        try:
            target.write_bytes(
                synthesize(args.engine, spoken, args.speaker, args.speed)
            )
        except requests.RequestException as e:
            failed.append(f"{fid}: {e}")
            continue

        ws.cell(r, header[AUDIO_COL]).value = filename
        made += 1

    if made:
        # 案内アプリ側で選べるように、この声を一覧へ登録する
        label = args.label
        if not label:
            speakers = requests.get(f"{args.engine}/speakers", timeout=10).json()
            for speaker in speakers:
                for style in speaker["styles"]:
                    if style["id"] == args.speaker:
                        label = f"{speaker['name']}・{style['name']}"
        register_set(wb, args.set, label or f"話者{args.speaker}",
                     args.speaker, args.speed)

    wb.save(path)

    print(f"\n作成: {made} 件 / 既にあるため飛ばした: {skipped} 件")
    for f in failed:
        print("  × " + f)
    if made:
        size = sum(f.stat().st_size for f in out_dir.glob("*.wav")) / 1024 / 1024
        print(f"音声の合計: {size:.1f} MB（{out_dir}）")
        print("\n管理画面の「案内アプリへ反映する」から書き出すと、案内アプリで使われます。")
        print("設置・公開で使う場合は、話者のクレジット表記を忘れずに。")


if __name__ == "__main__":
    main()
