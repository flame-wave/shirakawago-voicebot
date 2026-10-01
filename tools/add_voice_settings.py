"""読み上げの設定を、Excelから決められるようにする。

これまで声の設定は端末側（案内画面の ?setup=1）にしかなく、
職員が全台まわって設定する必要があった。
「どの端末でも同じにしたいこと」をここに移す。

【端末側に残るもの】
その端末に入っている声の一覧は、端末しか知らない。
Windowsに日本語とドイツ語、iPadに日本語とフランス語、という具合に違う。
管理者画面からは「どれを優先するか」までしか決められないので、
実際にどれを使うかの最終決定は端末側に残る。
"""
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PATH = ROOT / "tools" / "faq_master.xlsx"

SHEET = "読み上げ"
HEADERS = ["言語", "優先する声", "速さ", "高さ", "備考"]
WIDTHS = [10, 56, 8, 8, 40]

# 言語ごとの初期値。これまで js/config.js に書いてあったものを移す。
# 上から順に探し、その端末に入っている最初の声を使う。
ROWS = [
    ("ja", "Google 日本語, Microsoft Nanami, Microsoft Ayumi, Microsoft Haruka, "
           "Kyoko, O-ren", 0.95, 1.0,
     "Androidアプリで使っていたのと同じ「Google 日本語」を先頭にしている"),
    ("en", "Google US English, Microsoft Aria, Microsoft Jenny, Microsoft Michelle, "
           "Microsoft Zira, Samantha", 0.95, 1.0, ""),
    ("zh", "Google 普通话, Microsoft Xiaoxiao, Microsoft Xiaoyi, Microsoft Huihui, "
           "Tingting, Ting-Ting", 0.95, 1.0, ""),
    ("ko", "Google 한국의, Microsoft SunHi, Microsoft Heami, Yuna", 0.95, 1.0, ""),
    ("es", "Google español, Microsoft Elvira, Microsoft Helena, Microsoft Laura, "
           "Mónica, Monica, Paulina", 0.95, 1.0, ""),
    ("fr", "Google français, Microsoft Denise, Microsoft Vivienne, "
           "Microsoft Hortense, Amélie, Amelie, Audrey", 0.95, 1.0, ""),
]

GUIDE = (
    "※ 「優先する声」は上から順に探し、その端末に入っている最初の声を使います。"
    "名前の一部が合えば採用します（「Microsoft Aria - English (United States)」に"
    "「Microsoft Aria」で当たります）。"
    "どれも入っていない端末では、その言語の声から自動で選びます。"
    "速さは1.0が標準、小さいほどゆっくり。高さは1.0が標準、大きいほど高い声になります。"
)


def add_default_column(wb):
    """「音声セット」シートに、既定を選ぶ列を足す。"""
    ws = wb["音声セット"]
    header = {str(ws.cell(1, c).value).strip(): c
              for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    if "既定" in header:
        print("「音声セット」シートの「既定」列はすでにあります。")
        return

    at = ws.max_column + 1
    cell = ws.cell(row=1, column=at, value="既定")
    cell.font = Font(name="Yu Gothic", size=10, bold=True)
    cell.fill = PatternFill("solid", fgColor="EDF4E6")
    ws.column_dimensions[get_column_letter(at)].width = 8

    # 先頭の行を既定にしておく（これまでの動きと同じ）
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, header["名前"]).value
        if name and not str(name).startswith("※"):
            ws.cell(row=r, column=at, value="TRUE").font = Font(name="Yu Gothic", size=10)
            break
    print("「音声セット」シートに「既定」列を足しました。")


def add_voice_sheet(wb):
    if SHEET in wb.sheetnames:
        print(f"「{SHEET}」シートはすでにあります。中身は触りません。")
        return

    ws = wb.create_sheet(SHEET)
    for c, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(name="Yu Gothic", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="EDF4E6")
    ws.freeze_panes = "A2"
    for c, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width

    for r, values in enumerate(ROWS, start=2):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name="Yu Gothic", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (2, 5)))

    guide = ws.cell(row=len(ROWS) + 3, column=1, value=GUIDE)
    guide.font = Font(name="Yu Gothic", size=9, color="5A6B58")
    guide.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=len(ROWS) + 3, start_column=1,
                   end_row=len(ROWS) + 5, end_column=5)
    print(f"「{SHEET}」シートを作り、{len(ROWS)} 言語ぶんを入れました。")


def main():
    wb = load_workbook(PATH)
    add_default_column(wb)
    add_voice_sheet(wb)
    wb.save(PATH)


if __name__ == "__main__":
    main()
