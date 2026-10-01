"""「キャラクター」シートを作る。

案内画面に立つキャラクターを、職員が差し替え・追加できるようにする。
画像は assets/character/ に置き、このシートにファイル名を書く。

【表情の3枚が無くてもよい】
「通常」だけ書けば、聞き取り中も話している間も同じ絵を使う。
1枚しかないキャラクターでもそのまま動く。

【口の高さ・顔の広さ】
吹き出しの尻尾をキャラクターの口元に向けるために使う。
立ち絵ごとに顔の位置が違うので、同じ数字では合わない。
画像の上端から口まで、右端から顔の左端までを、高さに対する割合で書く。
空欄なら既定値（0.14 / 0.31）を使う。
"""
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PATH = ROOT / "tools" / "faq_master.xlsx"
SHEET = "キャラクター"

HEADERS = ["名前", "ID", "通常", "聞き取り中", "話している",
           "口の高さ", "顔の広さ", "既定", "有効", "備考"]
WIDTHS = [18, 14, 18, 18, 18, 10, 10, 8, 8, 40]

# 最初に入れておくもの。名前は仮なので、職員の方で付け替えてよい。
ROWS = [
    ("ナビ", "navi", "idle.png", "listening.png", "talking.png",
     0.14, 0.31, "TRUE", "TRUE",
     "表情が3枚ある案内役。これまで使っていたもの"),
    ("ゆい", "main", "main.png", "", "",
     0.16, 0.30, "", "TRUE",
     "白川郷のメイン。表情差分なし。名前は仮（「結」から）"),
    ("がっしょくん", "yuru", "yuru.png", "", "",
     0.40, 0.90, "", "TRUE",
     "ゆるキャラ。表情差分なし。名前は仮"),
    ("やねくん", "yane", "2.png", "", "",
     0.45, 0.90, "", "TRUE",
     "看板風。名前は仮。キャラクターとして使わないなら有効をFALSEに"),
    ("ごうちゃん", "gou", "4.png", "", "",
     0.45, 0.90, "", "TRUE",
     "看板風。名前は仮。キャラクターとして使わないなら有効をFALSEに"),
]


def main():
    wb = load_workbook(PATH)
    if SHEET in wb.sheetnames:
        print(f"「{SHEET}」シートはすでにあります。中身は触りません。")
        wb.close()
        return

    ws = wb.create_sheet(SHEET)
    for c, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(name="Yu Gothic", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="EDF4E6")
    ws.freeze_panes = "A2"
    for c, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width

    for r, values in enumerate(ROWS, start=2):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name="Yu Gothic", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(c == 10))

    guide = ws.cell(
        row=len(ROWS) + 3, column=1,
        value="※ 画像は assets/character/ に置く。「通常」だけ書けば1枚でも動く。"
              "「既定」にTRUEを入れた1件が、端末を開いたときのキャラクターになる。",
    )
    guide.font = Font(name="Yu Gothic", size=9, color="5A6B58")

    wb.save(PATH)
    print(f"「{SHEET}」シートを作り、{len(ROWS)} 件を入れました。")
    for row in ROWS:
        print(f"  {row[0]:<8} {row[1]:<6} {row[2]}")


if __name__ == "__main__":
    main()
