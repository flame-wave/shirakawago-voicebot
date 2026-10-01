"""
質問回答集に「言い換え」シートを用意する（初回のみ実行）。

観光客は、登録した質問例とは違う語で聞いてくる。
「銀行ありますか」「鞄を置きたい」「山に登れますか」——
語が違うだけで外れてしまうのを、この表で吸収する。

    python tools/init_synonyms.py --source tools/faq_master.xlsx

すでにシートがある場合は、足りない行だけ追加する（手で直した内容は消さない）。

【表の意味】
  代表語        … 質問例に実際に書かれている語
  同じ意味の語  … 観光客が言いそうな別の言い方（1行に1つ）

検索のときは、入力文に「同じ意味の語」が含まれていたら、
その「代表語」も入力文に含まれているものとして扱う。
"""

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET = "言い換え"
HEADERS = ["言語", "代表語", "同じ意味の語", "備考"]
FONT = "Yu Gothic"

# 初期の言い換え。語は「代表語が質問例に実際にある」ことを確かめて選んである。
ENTRIES = [
    ("ja", "atm", ["銀行", "キャッシュディスペンサー", "現金 引き出し"], "お金をおろす場所"),
    ("ja", "両替", ["ドル", "ユーロ", "外貨", "替えられ", "替えたい"], ""),
    ("ja", "小銭", ["百円玉", "硬貨", "コイン", "くずし", "崩し"], "ロッカー用の小銭"),
    ("ja", "荷物", ["鞄", "かばん", "カバン", "バッグ", "リュック", "旅行かばん"], ""),
    ("ja", "預け", ["置いて", "置いとき", "あずけ", "預かって"], "荷物を置く言い方"),
    ("ja", "駐車場", ["駐車スペース", "駐車場所", "パーキング", "停められ", "車を停め"], ""),
    ("ja", "料金", ["お金", "いくら", "値段", "費用", "有料"], ""),
    ("ja", "時間", ["どのくらい", "どれくらい", "どのぐらい", "どれぐらい", "何分"], ""),
    ("ja", "回る", ["回れ", "回り", "まわれ", "まわる", "巡り"], ""),
    ("ja", "見学", ["観光", "見られ", "見たい", "見て回", "中に入"], ""),
    ("ja", "車椅子", ["車いす", "車イス", "くるまいす"], ""),
    ("ja", "加熱式タバコ", ["電子タバコ", "電子たばこ", "アイコス", "ベイプ", "アイコス"], ""),
    ("ja", "喫煙所", ["吸え", "吸い", "吸っ", "喫煙", "スモーキング"], ""),
    ("ja", "住民", ["住んでいる", "住んでる", "地元の方", "暮らしている"], ""),
    ("ja", "マナー", ["気をつけ", "注意点", "守ること", "ルール", "いけないこと"], ""),
    ("ja", "雪遊び", ["そり", "ソリ", "雪だるま", "かまくら", "雪滑り"], ""),
    ("ja", "ごみ", ["捨て", "廃棄", "空き缶", "ペットボトル", "生ごみ"], ""),
    ("ja", "今日", ["飛び込み", "今から", "今夜", "今晩", "当日"], "当日の宿"),
    ("ja", "忘れ", ["置いてき", "置き忘れ", "残してき", "見当たらな"], ""),
    ("ja", "お手洗い", ["トイレ", "化粧室", "便所", "多目的トイレ"], ""),
    ("ja", "飲食店", ["食事", "食べ", "食堂", "ご飯屋", "ランチ"], ""),
    ("ja", "登山", ["山に登", "登れ", "山登り", "山歩き"], ""),
    ("ja", "早朝", ["朝早く", "朝一", "何時から開"], ""),
    ("ja", "タクシー", ["配車", "ハイヤー"], ""),
    ("ja", "戻る", ["帰り", "戻り", "帰る", "引き返"], ""),
    ("ja", "レンタル", ["借り", "貸し出し", "貸出し", "借りられ"], ""),
    ("ja", "宿", ["泊まれ", "泊まりたい", "宿泊", "民宿", "ホテル"], ""),
]


def ensure_sheet(wb):
    if SHEET in wb.sheetnames:
        return wb[SHEET], False
    ws = wb.create_sheet(SHEET)
    for c, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(name=FONT, size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="EDF4E6")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for c, width in enumerate([8, 18, 34, 28], start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    return ws, True


def existing_keys(ws):
    keys = set()
    for r in range(2, ws.max_row + 1):
        lang = (ws.cell(r, 1).value or "").strip().lower()
        rep = (ws.cell(r, 2).value or "").strip()
        if rep:
            keys.add((lang, rep))
    return keys


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", default="tools/faq_master.xlsx")
    args = ap.parse_args()

    path = Path(args.source)
    wb = load_workbook(path)
    ws, created = ensure_sheet(wb)
    have = existing_keys(ws)

    row = ws.max_row + 1 if not created else 2
    added = 0
    for lang, rep, words, note in ENTRIES:
        if (lang, rep) in have:
            continue  # 手で直したものを上書きしない
        ws.cell(row=row, column=1, value=lang)
        ws.cell(row=row, column=2, value=rep)
        ws.cell(row=row, column=3, value="\n".join(words))
        ws.cell(row=row, column=4, value=note or None)
        for c in range(1, 5):
            cell = ws.cell(row=row, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (3, 4)))
        row += 1
        added += 1

    wb.save(path)
    print(f"{'作成' if created else '追記'}しました: {SHEET} シート（{added} 行）")


if __name__ == "__main__":
    main()
